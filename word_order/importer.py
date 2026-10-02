import os
import re
import sqlite3
from datetime import date, datetime

from openpyxl import load_workbook

from word_order.storage import initialize_database


TABLES = {
    "tblLevels": {
        "destination": "levels",
        "primary_key": ("level_id",),
        "columns": (
            "level_id", "level_name_en", "level_name_es", "cefr",
            "grammar_focus", "unlocked_structures", "target_tokens", "notes_es",
        ),
        "required": ("level_id", "level_name_en", "level_name_es"),
    },
    "tblWords": {
        "destination": "words",
        "primary_key": ("word_id",),
        "columns": (
            "word_id", "english", "spanish", "pos", "subtype", "countability",
            "animacy", "verb_class", "cefr", "level_min", "tags", "active",
        ),
        "required": ("word_id", "english", "spanish", "pos", "level_min", "active"),
    },
    "tblVerbForms": {
        "destination": "verb_forms",
        "primary_key": ("base",),
        "columns": (
            "base", "third_person", "past", "past_participle", "ing", "spanish",
            "verb_class", "allows_iobj", "complement_type", "irregular", "level_min",
        ),
        "required": ("base", "spanish", "verb_class", "level_min"),
    },
    "tblPatterns": {
        "destination": "patterns",
        "primary_key": ("pattern_id",),
        "columns": (
            "pattern_id", "domain", "family", "pattern_name", "template",
            "spanish_guide", "min_level", "max_tokens", "recursive", "example_en",
        ),
        "required": ("pattern_id", "pattern_name", "template", "min_level"),
    },
    "tblPatternSlots": {
        "destination": "pattern_slots",
        "primary_key": ("pattern_id", "slot_order"),
        "columns": (
            "pattern_id", "slot_order", "slot_code", "required", "allowed_pos",
            "allowed_subtype", "agreement_rule", "notes_es",
        ),
        "required": ("pattern_id", "slot_order", "slot_code"),
    },
    "tblChallenges": {
        "destination": "challenges",
        "primary_key": ("challenge_id",),
        "columns": (
            "challenge_id", "level_id", "pattern_id", "prompt_es", "correct_en",
            "shuffled_tokens", "accept_alt_en", "explanation_es", "skill_tag",
            "difficulty", "active", "source_type", "token_count",
        ),
        "required": (
            "challenge_id", "level_id", "pattern_id", "prompt_es", "correct_en",
            "shuffled_tokens", "active",
        ),
    },
}

REQUIRED_HEADERS = {
    table_name: tuple(
        column for column in definition["columns"]
        if column != "challenge_count"
    )
    for table_name, definition in TABLES.items()
}
HEADER_OVERRIDES = {"third_person": "third_person"}
BOOLEAN_COLUMNS = {"active"}
INTEGER_COLUMNS = {
    "level_id", "level_min", "min_level", "max_tokens", "slot_order", "token_count",
}


class WorkbookImportError(ValueError):
    def __init__(self, message, rejected=0):
        super().__init__(message)
        self.rejected = rejected


def _column_name(header):
    normalized = re.sub(r"[^a-z0-9]+", "_", str(header).strip().casefold()).strip("_")
    return HEADER_OVERRIDES.get(normalized, normalized)


def _cell_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _as_integer(value, field, table_name):
    if value is None:
        return None
    if isinstance(value, bool):
        raise WorkbookImportError(f"{table_name}.{field} must be an integer.", rejected=1)
    try:
        converted = int(value)
    except (TypeError, ValueError, OverflowError) as error:
        raise WorkbookImportError(
            f"{table_name}.{field} has invalid value {value!r}.", rejected=1
        ) from error
    if str(value).strip() not in {str(converted), f"{converted}.0"} and not isinstance(value, int):
        raise WorkbookImportError(
            f"{table_name}.{field} has invalid value {value!r}.", rejected=1
        )
    return converted


def _as_active(value, table_name, row_number):
    normalized = str(value).strip().casefold()
    if normalized in {"y", "yes", "true", "1"}:
        return 1
    if normalized in {"n", "no", "false", "0"}:
        return 0
    raise WorkbookImportError(
        f"{table_name} row {row_number} has invalid Active value {value!r}.",
        rejected=1,
    )


def _read_workbook(path):
    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as error:
        raise WorkbookImportError(f"Could not open workbook: {error}") from error

    try:
        found_tables = {}
        for worksheet in workbook.worksheets:
            for table_name in worksheet.tables:
                if table_name in found_tables:
                    raise WorkbookImportError(f"Duplicate workbook table: {table_name}")
                found_tables[table_name] = (worksheet, worksheet.tables[table_name])

        missing_tables = sorted(set(TABLES) - set(found_tables))
        if missing_tables:
            raise WorkbookImportError(
                "Missing workbook tables: " + ", ".join(missing_tables)
            )

        records = {}
        for table_name, definition in TABLES.items():
            worksheet, table = found_tables[table_name]
            cells = worksheet[table.ref]
            headers = [_column_name(cell.value) if cell.value is not None else "" for cell in cells[0]]
            missing_headers = sorted(set(REQUIRED_HEADERS[table_name]) - set(headers))
            if missing_headers:
                raise WorkbookImportError(
                    f"{table_name} is missing columns: {', '.join(missing_headers)}"
                )

            table_records = []
            for row_number, row in enumerate(cells[1:], start=2):
                values = {
                    header: _cell_value(cell.value)
                    for header, cell in zip(headers, row)
                    if header
                }
                if not any(value is not None for value in values.values()):
                    continue
                record = {
                    column: values.get(column)
                    for column in definition["columns"]
                }
                if table_name == "tblLevels":
                    record.pop("challenge_count", None)
                for field in INTEGER_COLUMNS.intersection(record):
                    if record[field] is not None:
                        record[field] = _as_integer(record[field], field, table_name)
                if "active" in record:
                    record["active"] = _as_active(record["active"], table_name, row_number)

                missing_fields = [
                    field for field in definition["required"]
                    if record.get(field) is None
                ]
                if missing_fields:
                    raise WorkbookImportError(
                        f"{table_name} row {row_number} is missing: {', '.join(missing_fields)}",
                        rejected=1,
                    )
                table_records.append(record)
            records[table_name] = table_records
        return records
    finally:
        workbook.close()


def _validate_relationships(records):
    for table_name, definition in TABLES.items():
        seen = set()
        for index, record in enumerate(records[table_name], start=2):
            identity = tuple(record[key] for key in definition["primary_key"])
            if identity in seen:
                raise WorkbookImportError(
                    f"{table_name} contains duplicate ID {identity!r} at row {index}.",
                    rejected=1,
                )
            seen.add(identity)

    level_ids = {row["level_id"] for row in records["tblLevels"]}
    pattern_ids = {row["pattern_id"] for row in records["tblPatterns"]}
    for index, challenge in enumerate(records["tblChallenges"], start=2):
        if challenge["level_id"] not in level_ids:
            raise WorkbookImportError(
                f"tblChallenges row {index} references missing level {challenge['level_id']!r}.",
                rejected=1,
            )
        if challenge["pattern_id"] not in pattern_ids:
            raise WorkbookImportError(
                f"tblChallenges row {index} references missing pattern {challenge['pattern_id']!r}.",
                rejected=1,
            )
    for index, slot in enumerate(records["tblPatternSlots"], start=2):
        if slot["pattern_id"] not in pattern_ids:
            raise WorkbookImportError(
                f"tblPatternSlots row {index} references missing pattern {slot['pattern_id']!r}.",
                rejected=1,
            )


def import_workbook(workbook_path, database_path, deactivate_missing=False):
    workbook_path = os.path.abspath(os.fspath(workbook_path))
    database_path = os.path.abspath(os.fspath(database_path))
    if not os.path.isfile(workbook_path):
        raise WorkbookImportError(f"Workbook not found: {workbook_path}")

    records = _read_workbook(workbook_path)
    _validate_relationships(records)
    initialize_database(database_path)
    summary = {"inserted": 0, "updated": 0, "unchanged": 0, "deactivated": 0, "rejected": 0}
    connection = sqlite3.connect(database_path)
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        connection.execute("BEGIN IMMEDIATE")
        for table_name, definition in TABLES.items():
            destination = definition["destination"]
            columns = definition["columns"]
            primary_key = definition["primary_key"]
            placeholders = ", ".join("?" for _ in columns)
            column_sql = ", ".join(columns)
            update_sql = ", ".join(
                f"{column} = excluded.{column}"
                for column in columns if column not in primary_key
            )
            insert_sql = (
                f"INSERT INTO {destination} ({column_sql}) VALUES ({placeholders}) "
                f"ON CONFLICT ({', '.join(primary_key)}) DO UPDATE SET {update_sql}"
            )
            select_sql = (
                f"SELECT {column_sql} FROM {destination} WHERE "
                + " AND ".join(f"{key} = ?" for key in primary_key)
            )
            for record in records[table_name]:
                key = tuple(record[field] for field in primary_key)
                existing = connection.execute(select_sql, key).fetchone()
                values = tuple(record[column] for column in columns)
                if existing is None:
                    summary["inserted"] += 1
                elif tuple(existing) == values:
                    summary["unchanged"] += 1
                    continue
                else:
                    summary["updated"] += 1
                connection.execute(insert_sql, values)

            if deactivate_missing and destination in {"words", "challenges"}:
                active_ids = {
                    record[primary_key[0]] for record in records[table_name]
                }
                if active_ids:
                    placeholders = ", ".join("?" for _ in active_ids)
                    cursor = connection.execute(
                        f"UPDATE {destination} SET active = 0 WHERE active = 1 "
                        f"AND {primary_key[0]} NOT IN ({placeholders})",
                        tuple(active_ids),
                    )
                else:
                    cursor = connection.execute(
                        f"UPDATE {destination} SET active = 0 WHERE active = 1"
                    )
                summary["deactivated"] += cursor.rowcount

        stat = os.stat(workbook_path)
        connection.execute(
            "INSERT INTO word_order_source_state (source_path, modified_ns, file_size) "
            "VALUES (?, ?, ?) ON CONFLICT(source_path) DO UPDATE SET "
            "modified_ns = excluded.modified_ns, file_size = excluded.file_size",
            (workbook_path, stat.st_mtime_ns, stat.st_size),
        )
        connection.commit()
    except Exception as error:
        connection.rollback()
        if isinstance(error, WorkbookImportError):
            raise
        raise WorkbookImportError(f"Workbook import rolled back: {error}") from error
    finally:
        connection.close()
    return summary


def synchronize_if_changed(workbook_path, database_path):
    workbook_path = os.path.abspath(os.fspath(workbook_path))
    database_path = os.path.abspath(os.fspath(database_path))
    initialize_database(database_path)
    stat = os.stat(workbook_path)
    connection = sqlite3.connect(database_path)
    try:
        previous = connection.execute(
            "SELECT modified_ns, file_size FROM word_order_source_state WHERE source_path = ?",
            (workbook_path,),
        ).fetchone()
    finally:
        connection.close()
    if previous == (stat.st_mtime_ns, stat.st_size):
        return {"inserted": 0, "updated": 0, "unchanged": 0, "deactivated": 0, "rejected": 0}
    return import_workbook(workbook_path, database_path)