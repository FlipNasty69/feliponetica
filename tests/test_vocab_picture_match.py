import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from urllib.parse import unquote, urlsplit

from openpyxl import Workbook, load_workbook

from app import app
import games.vocab_game as vocab_game
from games.vocab_game import (
    _build_word,
    _labels_match,
    _matching_image_references,
    _reference_index,
    vocabulary_sets,
)


WORKBOOK_PATH = Path(__file__).resolve().parents[1] / "static" / "data" / "Vocab_Game.xlsx"


class VocabularyGameTestCase(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.client = app.test_client()

    def test_game_page_offers_numbered_sets_and_both_progressive_phases(self):
        response = self.client.get("/vocab/game")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'id="setSelect"', response.data)
        self.assertIn(b"Set 1.1", response.data)
        self.assertIn(b"Set 1.99", response.data)
        self.assertIn(b"no playable media", response.data)
        self.assertIn(b"Phase 1", response.data)
        self.assertIn(b"Words + audio", response.data)
        self.assertIn(b"Phase 2", response.data)
        self.assertIn(b"Audio only", response.data)
        self.assertIn(b"Match the English word", response.data)
        self.assertIn(b"Match the audio", response.data)
        self.assertIn(b"batchSize = 5", response.data)
        self.assertIn(b"itemsInCurrentRound", response.data)
        self.assertIn(
            b"vocabItems.slice(0, Math.min((batchIndex + 1) * batchSize, vocabItems.length))",
            response.data,
        )
        self.assertIn(b"Next five words", response.data)
        self.assertIn(b"learnAudioButton.onclick = () => playAudio(item.audio)", response.data)

    def test_every_set_sheet_column_is_exposed_in_workbook_order(self):
        sets = vocabulary_sets()
        workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
        try:
            sheet_set_ids = []
            for row in workbook["sets"].iter_rows(values_only=True):
                if not row or not str(row[0] or "").strip().casefold().startswith("set "):
                    continue
                sheet_set_ids.extend(
                    str(value).strip()
                    for value in row[1:]
                    if value is not None and str(value).strip()
                )
        finally:
            workbook.close()

        self.assertEqual(list(sets), sheet_set_ids)

        for set_id, vocabulary_set in sets.items():
            with self.subTest(set_id=set_id):
                words = vocabulary_set["words"]
                self.assertEqual(len({item["id"] for item in words}), len(words))
                for item in words:
                    self.assertTrue(item["audio"])
                    self.assertEqual(len(item["images"]), 1)

    def test_one_vocabulary_word_uses_one_image_when_multiple_images_match(self):
        image_references = [
            {"label": "bathroom sink", "order": 0},
            {"label": "kitchen sink", "order": 1},
        ]

        matches = _matching_image_references(
            "sink",
            image_references,
            _reference_index(image_references, include_suffix=True),
        )

        self.assertEqual(len(matches), 1)

    def test_word_builder_does_not_combine_images_from_multiple_rows(self):
        item = _build_word(
            "1.2",
            "sink",
            [
                {
                    "audio": "/static/audio/vocab_game/sink.opus",
                    "images": ["/static/images/vocab_game/bathroom_sink.webp"],
                },
                {
                    "audio": "/static/audio/vocab_game/sink.opus",
                    "images": ["/static/images/vocab_game/kitchen_sink.webp"],
                },
            ],
            0,
        )

        self.assertEqual(
            item["images"],
            ["/static/images/vocab_game/bathroom_sink.webp"],
        )

    def test_sets_api_returns_all_numbered_sets_and_playable_word_counts(self):
        response = self.client.get("/vocab/api/get-vocab-sets")

        self.assertEqual(response.status_code, 200)
        api_sets = response.get_json()
        sets_by_id = {item["id"]: item for item in api_sets}
        self.assertEqual(list(sets_by_id), list(vocabulary_sets()))
        self.assertIn("2.2", sets_by_id)
        self.assertGreater(sets_by_id["2.2"]["word_count"], 0)
        self.assertIn("1.41", sets_by_id)
        self.assertIn("1.13", sets_by_id)
        self.assertEqual(sets_by_id["1.13"]["word_count"], 0)

    def test_sink_phrases_use_their_workbook_images_and_audio(self):
        bathroom_sink = next(
            item
            for vocabulary_set in vocabulary_sets().values()
            for item in vocabulary_set["words"]
            if item["english"].casefold() == "bathroom sink"
        )
        kitchen_sink = next(
            item
            for vocabulary_set in vocabulary_sets().values()
            for item in vocabulary_set["words"]
            if item["english"].casefold() == "kitchen sink"
        )

        self.assertTrue(
            unquote(urlsplit(bathroom_sink["audio"]).path).endswith(
                "/bathroom sink.opus"
            )
        )
        self.assertEqual(
            {urlsplit(image).path.rsplit("/", 1)[-1] for image in bathroom_sink["images"]},
            {"bathroom_sink.webp"},
        )
        self.assertEqual(len(bathroom_sink["images"]), 1)
        self.assertTrue(
            unquote(urlsplit(kitchen_sink["audio"]).path).endswith(
                "/kitchen sink.opus"
            )
        )
        self.assertEqual(
            {urlsplit(image).path.rsplit("/", 1)[-1] for image in kitchen_sink["images"]},
            {"kitchen_sink.webp"},
        )
        self.assertEqual(len(kitchen_sink["images"]), 1)
        self.assertTrue(_labels_match("sing", "sink"))

    def test_numbered_sets_use_their_workbook_sections(self):
        plant_set_words = {
            item["english"].casefold()
            for item in vocabulary_sets()["3.93"]["words"]
        }
        profession_set_words = {
            item["english"].casefold()
            for item in vocabulary_sets()["1.41"]["words"]
        }

        self.assertIn("plant", plant_set_words)
        self.assertIn("cook", profession_set_words)
        self.assertIn("doctor", profession_set_words)

    def test_workbook_accepts_image_reference_sheet_name(self):
        workbook = Workbook()
        sets_sheet = workbook.active
        sets_sheet.title = "sets"
        sets_sheet.append(("set 1", "1.2"))
        sets_sheet.append(("word1", "sink"))

        audio_sheet = workbook.create_sheet("Audio Reference")
        audio_sheet.append(("word", "audio URL"))
        audio_sheet.append(
            ("sink", "https://cdn.example.test/audio/vocab_game/sink.opus")
        )

        image_sheet = workbook.create_sheet("Image Reference")
        image_sheet.append(("word", "image URL"))
        image_sheet.append(
            ("sink", "https://cdn.example.test/images/vocab_game/sink.webp")
        )

        with TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "vocabulary.xlsx"
            workbook.save(workbook_path)
            workbook.close()

            with patch.object(vocab_game, "VOCAB_WORKBOOK", workbook_path):
                vocabulary_sets.cache_clear()
                try:
                    result = vocabulary_sets()
                finally:
                    vocabulary_sets.cache_clear()

        sink = result["1.2"]["words"][0]
        self.assertEqual(sink["english"], "sink")
        self.assertEqual(len(sink["images"]), 1)
        self.assertTrue(sink["audio"].endswith("/sink.opus"))

    def test_selected_set_api_returns_playable_linked_media(self):
        response = self.client.get("/vocab/api/get-set-details?id=3.2-kitchen")

        self.assertEqual(response.status_code, 200)
        words = response.get_json()
        self.assertGreater(len(words), 5)
        self.assertTrue(all(item["audio"] and item["images"] for item in words))

        for item in words:
            for url in item["images"] + [item["audio"]]:
                path = urlsplit(url).path
                if path.startswith("/static/"):
                    asset_response = self.client.get(path)
                    self.assertEqual(asset_response.status_code, 200, url)
                    asset_response.close()

    def test_unknown_set_returns_not_found(self):
        response = self.client.get("/vocab/api/get-set-details?id=not-a-set")
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.get_json(), {"error": "Set not found"})

    def test_set_without_playable_media_is_known_but_cannot_start(self):
        response = self.client.get("/vocab/api/get-set-details?id=1.13")
        self.assertEqual(response.status_code, 422)
        self.assertIn("playable image/audio pairs", response.get_json()["error"])

    def test_legacy_vocabulary_routes_redirect_to_the_single_game(self):
        for path in (
            "/vocab/picture-match",
            "/vocab/picture-match/along",
            "/vocab/picture-match/french-audio",
            "/vocab/along",
        ):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 302)
                self.assertEqual(response.headers["Location"], "/vocab/game")

    def test_dashboard_has_one_vocabulary_game_entry(self):
        dashboard = self.client.get("/games")

        self.assertEqual(dashboard.status_code, 200)
        self.assertIn(b'href="/vocab/game"', dashboard.data)
        self.assertIn(b"Vocabulary Drill", dashboard.data)
        self.assertNotIn(b"Picture Match Vocabulary", dashboard.data)
        self.assertNotIn(b"Along Picture Match", dashboard.data)
        self.assertNotIn(b"French Audio and Picture Match", dashboard.data)


if __name__ == "__main__":
    unittest.main()
