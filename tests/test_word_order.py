import shutil
import sqlite3
import tempfile
import unittest
import warnings
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils import get_column_letter

from app import app
from word_order.gameplay import join_tokens, normalize_answer, split_tokens
from word_order.importer import WorkbookImportError, import_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "static" / "data" / "muster_structure.xlsx"


class WordOrderTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.database_path = Path(self.temp_dir.name) / "word_order.sqlite3"
        self.previous_database = app.config.get("WORD_ORDER_DATABASE")
        self.previous_workbook = app.config.get("WORD_ORDER_WORKBOOK")
        self.previous_testing = app.config.get("TESTING")
        app.config.update(
            TESTING=True,
            WORD_ORDER_DATABASE=str(self.database_path),
            WORD_ORDER_WORKBOOK=str(WORKBOOK),
        )
        self.client = app.test_client()

    def tearDown(self):
        if self.previous_database is None:
            app.config.pop("WORD_ORDER_DATABASE", None)
        else:
            app.config["WORD_ORDER_DATABASE"] = self.previous_database
        if self.previous_workbook is None:
            app.config.pop("WORD_ORDER_WORKBOOK", None)
        else:
            app.config["WORD_ORDER_WORKBOOK"] = self.previous_workbook
        app.config["TESTING"] = self.previous_testing
        self.temp_dir.cleanup()

    def connect(self):
        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    def test_blueprint_pages_and_dynamic_levels(self):
        page = self.client.get("/games/word-order/levels")
        levels = self.client.get("/games/word-order/api/levels").get_json()

        self.assertEqual(page.status_code, 200)
        self.assertEqual(len(levels), 12)
        self.assertEqual(levels[0]["status"], "available")
        self.assertEqual(levels[1]["status"], "locked")
        self.assertEqual(levels[-1]["name_en"], "Master structures")
        self.assertEqual(levels[-1]["challenge_count"], 6)

    def test_named_table_import_and_idempotent_upsert(self):
        first = import_workbook(WORKBOOK, self.database_path)
        second = import_workbook(WORKBOOK, self.database_path)
        connection = self.connect()
        try:
            counts = {
                table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                for table in ("levels", "words", "verb_forms", "patterns", "pattern_slots", "challenges")
            }
            foreign_key_errors = connection.execute("PRAGMA foreign_key_check").fetchall()
        finally:
            connection.close()

        self.assertEqual(first["inserted"], 597)
        self.assertEqual(second["unchanged"], 597)
        self.assertEqual(second["updated"], 0)
        self.assertEqual(counts, {
            "levels": 12, "words": 244, "verb_forms": 40,
            "patterns": 59, "pattern_slots": 170, "challenges": 72,
        })
        self.assertEqual(foreign_key_errors, [])

    def test_invalid_foreign_key_rolls_back_import(self):
        modified_workbook = Path(self.temp_dir.name) / "invalid.xlsx"
        shutil.copyfile(WORKBOOK, modified_workbook)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(modified_workbook)
            sheet = next(sheet for sheet in workbook.worksheets if "tblChallenges" in sheet.tables)
            table = sheet.tables["tblChallenges"]
            start_column, header_row, end_column, _ = range_boundaries(table.ref)
            headers = [
                sheet.cell(row=header_row, column=column).value
                for column in range(start_column, end_column + 1)
            ]
            pattern_column = headers.index("Pattern_ID") + start_column
            sheet.cell(row=header_row + 1, column=pattern_column).value = "P-MISSING"
            workbook.save(modified_workbook)
            workbook.close()

        with self.assertRaises(WorkbookImportError):
            import_workbook(modified_workbook, self.database_path)
        self.assertFalse(self.database_path.exists())

    def test_new_challenge_rows_are_imported_and_missing_rows_can_deactivate(self):
        import_workbook(WORKBOOK, self.database_path)
        modified_workbook = Path(self.temp_dir.name) / "updated.xlsx"
        shutil.copyfile(WORKBOOK, modified_workbook)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(modified_workbook)
            sheet = next(sheet for sheet in workbook.worksheets if "tblChallenges" in sheet.tables)
            table = sheet.tables["tblChallenges"]
            min_column, min_row, max_column, max_row = range_boundaries(table.ref)
            headers = [
                sheet.cell(row=min_row, column=column).value
                for column in range(min_column, max_column + 1)
            ]
            new_values = {
                "Challenge_ID": "C-ADDED",
                "Level_ID": 1,
                "Pattern_ID": "P001",
                "Prompt_ES": "ella",
                "Correct_EN": "She",
                "Shuffled_Tokens": "She",
                "Explanation_ES": "Pronombre sujeto.",
                "Skill_Tag": "pronoun",
                "Difficulty": "Easy",
                "Active": "Y",
                "Source_Type": "authored",
                "Token_Count": 1,
            }
            sheet.append([None] * max_column)
            appended_row = sheet.max_row
            for column, header in enumerate(headers, start=min_column):
                if header in new_values:
                    sheet.cell(row=appended_row, column=column).value = new_values[header]
            table.ref = f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{appended_row}"
            workbook.save(modified_workbook)
            workbook.close()

        app.config["WORD_ORDER_WORKBOOK"] = str(modified_workbook)
        levels = self.client.get("/games/word-order/api/levels").get_json()
        self.assertEqual(levels[0]["challenge_count"], 7)
        import_workbook(WORKBOOK, self.database_path, deactivate_missing=True)
        connection = self.connect()
        try:
            added_active = connection.execute(
                "SELECT active FROM challenges WHERE challenge_id = 'C-ADDED'"
            ).fetchone()[0]
            inactive_count = connection.execute(
                "SELECT COUNT(*) FROM challenges WHERE active = 0"
            ).fetchone()[0]
        finally:
            connection.close()
        self.assertEqual(added_active, 0)
        self.assertEqual(inactive_count, 1)

    def test_added_level_appears_without_a_fixed_level_limit(self):
        modified_workbook = Path(self.temp_dir.name) / "level-13.xlsx"
        shutil.copyfile(WORKBOOK, modified_workbook)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(modified_workbook)
            level_sheet = next(sheet for sheet in workbook.worksheets if "tblLevels" in sheet.tables)
            level_table = level_sheet.tables["tblLevels"]
            min_column, min_row, max_column, max_row = range_boundaries(level_table.ref)
            level_headers = [
                level_sheet.cell(row=min_row, column=column).value
                for column in range(min_column, max_column + 1)
            ]
            new_level = {
                "Level_ID": 13,
                "Level_Name_EN": "New content",
                "Level_Name_ES": "Contenido nuevo",
                "CEFR": "C1",
                "Grammar_Focus": "New grammar",
                "Unlocked_Structures": "EXTENSION",
                "Target_Tokens": 3,
                "Notes_ES": "Nivel agregado desde la hoja de cálculo.",
                "Challenge_Count": 1,
            }
            level_sheet.append([None] * max_column)
            appended_level_row = level_sheet.max_row
            for column, header in enumerate(level_headers, start=min_column):
                if header in new_level:
                    level_sheet.cell(row=appended_level_row, column=column).value = new_level[header]
            level_table.ref = f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{appended_level_row}"

            challenge_sheet = next(sheet for sheet in workbook.worksheets if "tblChallenges" in sheet.tables)
            challenge_table = challenge_sheet.tables["tblChallenges"]
            challenge_min_column, challenge_min_row, challenge_max_column, challenge_max_row = range_boundaries(challenge_table.ref)
            challenge_headers = [
                challenge_sheet.cell(row=challenge_min_row, column=column).value
                for column in range(challenge_min_column, challenge_max_column + 1)
            ]
            new_challenge = {
                "Challenge_ID": "C-LEVEL-13",
                "Level_ID": 13,
                "Pattern_ID": "P001",
                "Prompt_ES": "ella",
                "Correct_EN": "She",
                "Shuffled_Tokens": "She",
                "Explanation_ES": "Pronombre sujeto.",
                "Skill_Tag": "pronoun",
                "Difficulty": "Easy",
                "Active": "Y",
                "Source_Type": "authored",
                "Token_Count": 1,
            }
            challenge_sheet.append([None] * challenge_max_column)
            appended_challenge_row = challenge_sheet.max_row
            for column, header in enumerate(challenge_headers, start=challenge_min_column):
                if header in new_challenge:
                    challenge_sheet.cell(row=appended_challenge_row, column=column).value = new_challenge[header]
            challenge_table.ref = f"{get_column_letter(challenge_min_column)}{challenge_min_row}:{get_column_letter(challenge_max_column)}{appended_challenge_row}"
            workbook.save(modified_workbook)
            workbook.close()

        app.config["WORD_ORDER_WORKBOOK"] = str(modified_workbook)
        levels = self.client.get("/games/word-order/api/levels").get_json()
        self.assertEqual(len(levels), 13)
        self.assertEqual(levels[-1]["id"], 13)
        self.assertEqual(levels[-1]["challenge_count"], 1)
        self.assertEqual(levels[-1]["status"], "locked")

    def test_sync_cli_imports_workbook(self):
        runner = app.test_cli_runner()
        result = runner.invoke(args=["word-order", "sync-data", str(WORKBOOK)])
        self.assertEqual(result.exit_code, 0, result.output)
        self.assertIn("inserted=597", result.output)

    def test_level_route_starts_game_and_answers_are_server_checked(self):
        redirect = self.client.get("/games/word-order/level/1")
        self.assertEqual(redirect.status_code, 302)
        self.assertTrue(redirect.location.endswith("/games/word-order/play/1"))

        response = self.client.post("/games/word-order/api/session", json={"level_id": 1})
        state = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("correct_en", state)
        self.assertTrue(state["tokens"])
        resumed_before_check = self.client.get(
            f"/games/word-order/api/session/{state['session_id']}"
        ).get_json()
        self.assertEqual(resumed_before_check["tokens"], state["tokens"])

        connection = self.connect()
        try:
            correct_answer = connection.execute(
                "SELECT correct_en FROM challenges WHERE challenge_id = ?",
                (state["challenge_id"],),
            ).fetchone()[0]
        finally:
            connection.close()

        checked = self.client.post("/games/word-order/check", json={
            "session_id": state["session_id"], "answer": correct_answer,
        })
        self.assertTrue(checked.get_json()["correct"])
        resumed = self.client.get(
            f"/games/word-order/api/session/{state['session_id']}"
        ).get_json()
        self.assertTrue(resumed["checked"])
        self.assertTrue(resumed["feedback"]["correct"])
        duplicate = self.client.post("/games/word-order/check", json={
            "session_id": state["session_id"], "answer": correct_answer,
        })
        self.assertEqual(duplicate.status_code, 409)

    def test_eighty_percent_completion_unlocks_next_level(self):
        state = self.client.post(
            "/games/word-order/api/session", json={"level_id": 1}
        ).get_json()
        total = state["total_challenges"]
        for index in range(total):
            connection = self.connect()
            try:
                answer = connection.execute(
                    "SELECT correct_en FROM challenges WHERE challenge_id = ?",
                    (state["challenge_id"],),
                ).fetchone()[0]
            finally:
                connection.close()
            self.client.post("/games/word-order/check", json={
                "session_id": state["session_id"],
                "answer": answer if index < 5 else "not the answer",
            })
            advanced = self.client.post("/games/word-order/api/continue", json={
                "session_id": state["session_id"],
            }).get_json()
            if advanced["completed"]:
                break
            state = advanced["challenge"]

        levels = self.client.get("/games/word-order/api/levels").get_json()
        self.assertEqual(levels[0]["completion_percentage"], 83)
        self.assertEqual(levels[1]["status"], "available")

    def test_empty_level_returns_friendly_error(self):
        self.client.get("/games/word-order/api/levels")
        connection = self.connect()
        try:
            connection.execute("UPDATE challenges SET active = 0 WHERE level_id = 1")
            connection.commit()
        finally:
            connection.close()

        response = self.client.post(
            "/games/word-order/api/session", json={"level_id": 1}
        )
        self.assertEqual(response.status_code, 404)
        self.assertIn("desafíos activos", response.get_json()["error"])

    def test_token_handling_and_answer_normalization(self):
        self.assertEqual(split_tokens("to|go to|to"), ["to", "go to", "to"])
        self.assertEqual(join_tokens(["She", "will", "go", "?"]), "She will go?")
        self.assertEqual(normalize_answer("  She   will go ? "), "she will go?")
        self.assertEqual(normalize_answer("  Students  "), "students")


if __name__ == "__main__":
    unittest.main()