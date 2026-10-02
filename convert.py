import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent
SOURCE_FILE = BASE_DIR / "static" / "data" / "verb_conjugation.json"
OUTPUT_FILE = BASE_DIR / "verbs_master_table.xlsx"


def main():
    with SOURCE_FILE.open(encoding="utf-8") as source:
        data = json.load(source)

    rows = []
    verb_fields = []
    for parent_category, parent_data in data.items():
        for series in parent_data.get("series", []):
            for verb in series.get("verbs", []):
                rows.append({
                    "parent_category": parent_category,
                    "series_id": series.get("id", ""),
                    "series_title": series.get("title", ""),
                    **verb,
                })
                for field in verb:
                    if field not in verb_fields:
                        verb_fields.append(field)

    headers = ["parent_category", "series_id", "series_title", *verb_fields]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Verbs"
    worksheet.append(headers)

    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_index, header in enumerate(headers, start=1):
        values = [str(row.get(header, "")) for row in rows]
        width = min(max(len(header), *(len(value) for value in values)) + 2, 50)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    table = Table(displayName="VerbData", ref=worksheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)
    workbook.save(OUTPUT_FILE)
    print(f"Created {OUTPUT_FILE} with {len(rows)} verbs and {len(headers)} columns.")


if __name__ == "__main__":
    main()