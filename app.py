import sqlite3
import os
import smtplib
import json
import re
import secrets
import zipfile

from contextlib import closing
from email.message import EmailMessage
from flask import request, render_template
from flask import Flask, render_template, request, send_from_directory, abort
from flask import session, redirect, url_for
from features.feliponetica import feliponetica_bp
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from werkzeug.security import check_password_hash, generate_password_hash
from word_order.storage import connect_database as connect_word_order_database

app = Flask(__name__)

app.secret_key = os.environ.get("FLASK_SECRET_KEY") or secrets.token_hex(32)
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get(
        "FLASK_COOKIE_SECURE",
        "1" if os.environ.get("RENDER") else "0",
    ) == "1",
)
STUDENT_DATABASE_PATH = os.path.join(app.root_path, "student_users.db")
TEST_RESULTS_ADMIN_PASSWORD_HASH = os.environ.get(
    "TEST_RESULTS_ADMIN_PASSWORD_HASH",
    "pbkdf2:sha256:600000$pgzAbeZ0AQmCtlwz$249a511824fc4651e9f4a0be139f920088c6ab99db794a1ce50a8c45391107a8",
)
USERS = {
    "felipe": {
        "password": "felipe",
        "role": "admin",
    }
}

# ============================================================
# GAME BLUEPRINTS
# ============================================================
# Add future game blueprints here and then register them below.
# Example:
# from games.vocab_game import vocab_bp
# GAME_BLUEPRINTS = [vocab_bp]
from games.vocab_game import vocab_bp
from games.phantom_pronouns import phantom_pronouns_bp
from games.adverb_game import adverb_game_bp
from word_order import word_order_bp
from word_order.cli import register_cli
GAME_BLUEPRINTS = [vocab_bp, phantom_pronouns_bp, adverb_game_bp, word_order_bp]

def register_game_blueprints():
    for blueprint in GAME_BLUEPRINTS:
        app.register_blueprint(blueprint)


register_game_blueprints()
app.register_blueprint(feliponetica_bp)
register_cli(app)


@app.route("/static/data/tests.xlsx")
def private_test_workbook():
    abort(404)


class StudentDatabaseConnection:
    def __init__(self, connection):
        self.connection = connection

    def execute(self, query, parameters=()):
        return self.connection.execute(query, parameters)

    def executemany(self, query, parameters):
        return self.connection.executemany(query, parameters)

    def commit(self):
        self.connection.commit()

    def close(self):
        self.connection.close()

    def __enter__(self):
        self.connection.__enter__()
        return self

    def __exit__(self, exception_type, exception, traceback):
        return self.connection.__exit__(exception_type, exception, traceback)


def connect_student_database():
    connection = sqlite3.connect(prepare_database_path(student_database_path()))
    connection.row_factory = sqlite3.Row
    return StudentDatabaseConnection(connection)


def connect_test_results_database():
    if test_results_database_path() == student_database_path():
        return connect_student_database()
    connection = sqlite3.connect(prepare_database_path(test_results_database_path()))
    connection.row_factory = sqlite3.Row
    return StudentDatabaseConnection(connection)


def database_integrity_errors():
    return (sqlite3.IntegrityError,)


def student_database_path():
    return os.fspath(
        app.config.get("STUDENT_DATABASE_PATH")
        or os.environ.get("STUDENT_DATABASE_PATH")
        or STUDENT_DATABASE_PATH
    )


def prepare_database_path(database_path):
    database_path = os.path.abspath(os.fspath(database_path))
    os.makedirs(os.path.dirname(database_path), exist_ok=True)
    return database_path


def init_student_database():
    connection = connect_student_database()
    with closing(connection) as connection, connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS student_users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                last_name TEXT NOT NULL,
                username TEXT UNIQUE,
                email TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'student'
            )
        """)
        connection.execute("""
            CREATE TABLE IF NOT EXISTS student_test_approvals (
                user_id INTEGER NOT NULL,
                category TEXT NOT NULL,
                approved INTEGER NOT NULL DEFAULT 0,
                updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
                PRIMARY KEY (user_id, category),
                FOREIGN KEY (user_id) REFERENCES student_users(id)
            )
        """)
        connection.execute("""
            CREATE TABLE IF NOT EXISTS conjugation_progress (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                group_name TEXT NOT NULL,
                completed_series INTEGER NOT NULL DEFAULT 0,
                current_verb INTEGER NOT NULL DEFAULT 0,
                score INTEGER NOT NULL DEFAULT 0,
                UNIQUE(user_id, group_name),
                FOREIGN KEY (user_id) REFERENCES student_users(id)
            )
        """)
        columns = {
            row["name"]
            for row in connection.execute("PRAGMA table_info(conjugation_progress)").fetchall()
        }
        if "current_verb" not in columns:
            connection.execute(
                "ALTER TABLE conjugation_progress ADD COLUMN current_verb "
                "INTEGER NOT NULL DEFAULT 0"
            )
        student_columns = {
            row["name"]
            for row in connection.execute("PRAGMA table_info(student_users)").fetchall()
        }
        if "username" not in student_columns:
            connection.execute("ALTER TABLE student_users ADD COLUMN username TEXT")
        if "role" not in student_columns:
            connection.execute(
                "ALTER TABLE student_users ADD COLUMN role TEXT NOT NULL DEFAULT 'student'"
            )
        connection.execute(
            "CREATE UNIQUE INDEX IF NOT EXISTS student_users_username_idx ON student_users(username)"
        )
        for student in connection.execute(
            "SELECT id, password_hash FROM student_users"
        ).fetchall():
            user_id = student["id"]
            password_hash = student["password_hash"]
            if password_hash and not password_hash.startswith(
                ("scrypt:", "pbkdf2:", "argon2:")
            ):
                connection.execute(
                    "UPDATE student_users SET password_hash = ? WHERE id = ?",
                    (generate_password_hash(password_hash), user_id),
                )


init_student_database()


def test_results_database_path():
    return os.fspath(
        app.config.get("TEST_RESULTS_DATABASE_PATH")
        or os.environ.get("TEST_RESULTS_DATABASE_PATH")
        or os.path.join(app.instance_path, "test_results.sqlite3")
    )


def init_test_results_database():
    database_path = prepare_database_path(test_results_database_path())
    sqlite_connection = sqlite3.connect(database_path)
    sqlite_connection.row_factory = sqlite3.Row
    connection = StudentDatabaseConnection(sqlite_connection)
    with closing(connection) as connection, connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS test_results (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                test_name TEXT NOT NULL,
                name TEXT NOT NULL,
                email TEXT NOT NULL,
                score INTEGER NOT NULL,
                question_count INTEGER NOT NULL,
                report_json TEXT NOT NULL,
                created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """)
        columns = {
            row["name"] for row in connection.execute("PRAGMA table_info(test_results)")
        }
        if "user_id" not in columns:
            connection.execute("ALTER TABLE test_results ADD COLUMN user_id INTEGER")

    legacy_database_path = os.path.abspath(student_database_path())
    if (
        os.path.abspath(database_path) != legacy_database_path
        and os.path.isfile(legacy_database_path)
    ):
        with closing(sqlite3.connect(legacy_database_path)) as legacy_connection:
            legacy_connection.row_factory = sqlite3.Row
            legacy_tables = {
                row["name"]
                for row in legacy_connection.execute(
                    "SELECT name FROM sqlite_master WHERE type = 'table'"
                )
            }
            if "test_results" in legacy_tables:
                legacy_columns = {
                    row["name"]
                    for row in legacy_connection.execute(
                        "PRAGMA table_info(test_results)"
                    )
                }
                result_columns = [
                    "id",
                    "user_id",
                    "test_name",
                    "name",
                    "email",
                    "score",
                    "question_count",
                    "report_json",
                    "created_at",
                ]
                required_legacy_columns = set(result_columns) - {"user_id"}
                if required_legacy_columns.issubset(legacy_columns):
                    user_id_expression = (
                        "user_id" if "user_id" in legacy_columns else "NULL AS user_id"
                    )
                    legacy_results = legacy_connection.execute(
                        f"""
                        SELECT id, {user_id_expression}, test_name, name, email,
                            score, question_count, report_json, created_at
                        FROM test_results
                        """
                    ).fetchall()
                    if legacy_results:
                        placeholders = ", ".join("?" for _ in result_columns)
                        with closing(connect_test_results_database()) as results_connection, results_connection:
                            results_connection.executemany(
                                f"""
                                INSERT OR IGNORE INTO test_results (
                                    {", ".join(result_columns)}
                                ) VALUES ({placeholders})
                                """,
                                [tuple(result[column] for column in result_columns)
                                 for result in legacy_results],
                            )


init_test_results_database()


def _normalize_test_header(value):
    header = str(value or "").strip().lower()
    normalized = re.sub(r"[^a-z0-9]", "", header)
    if header.endswith("#") and normalized.startswith("question"):
        return "questionnumber"
    return normalized


def _test_header_indexes(header):
    aliases = {
        "id": {"id", "questionid", "questionnumber", "number", "qid"},
        "type": {"type", "questiontype"},
        "category": {"category", "topic", "section", "grammar"},
        "question": {"question", "questiontext", "prompt", "statement"},
        "choice_a": {"a", "choicea", "optiona", "answera"},
        "choice_b": {"b", "choiceb", "optionb", "answerb"},
        "choice_c": {"c", "choicec", "optionc", "answerc"},
        "choice_d": {"d", "choiced", "optiond", "answerd"},
        "correct_answer": {
            "correct",
            "correctanswer",
            "correctoption",
            "answer",
            "answerkey",
            "key",
        },
    }
    normalized = [_normalize_test_header(value) for value in header]
    indexes = {}
    for field, field_aliases in aliases.items():
        for index, value in enumerate(normalized):
            if value in field_aliases:
                indexes[field] = index
                break
    return indexes


def _parse_test_questions(rows, source_name, header=None, allow_legacy_layout=True):
    header_indexes = _test_header_indexes(header or ())
    uses_headers = (
        "question" in header_indexes
        and "choice_a" in header_indexes
        and "choice_b" in header_indexes
        and "correct_answer" in header_indexes
    )
    if header is not None and not uses_headers and (
        not allow_legacy_layout or len(header) < 11
    ):
        return []

    questions = []
    question_ids = set()
    for row in rows:
        if uses_headers:
            value = lambda field, fallback: (
                row[header_indexes[field]]
                if field in header_indexes and header_indexes[field] < len(row)
                else fallback
            )
            question_text = value("question", None)
            choices = {
                label: str(value(f"choice_{label.lower()}", None)).strip()
                for label in "ABCD"
                if value(f"choice_{label.lower()}", None) is not None
                and str(value(f"choice_{label.lower()}", None)).strip()
            }
            question_id_value = value("id", len(questions) + 1)
            question_type = value("type", "Multiple Choice")
            category = value("category", "")
            correct_answer_value = value("correct_answer", "")
        else:
            if len(row) < 11:
                continue
            question_text = row[5]
            if not question_text:
                continue
            question_id_value = row[0] or len(questions) + 1
            question_type = row[1] or "Multiple Choice"
            category = row[3] or ""
            choices = {
                label: str(row[column]).strip()
                for label, column in zip("ABCD", range(6, 10))
                if row[column] is not None and str(row[column]).strip()
            }
            correct_answer_value = row[10] or ""

        if len(row) >= 11 and row[5]:
            question_text = row[5]

        if not question_text:
            continue
        normalized_type = str(question_type or "").strip().lower()
        if normalized_type not in {"multiple choice", "multiple-choice", "mcq", "mc"}:
            raise ValueError(
                f"Unsupported question type in {source_name}, question "
                f"{question_id_value}: {question_type}"
            )
        question_id = str(question_id_value or len(questions) + 1).strip()
        correct_answer = str(correct_answer_value or "").strip().upper()
        if (
            question_id in question_ids
            or len(choices) < 2
            or correct_answer not in choices
        ):
            raise ValueError(
                f"Question {question_id} in {source_name} has a duplicate ID, "
                "incomplete choices, or an invalid answer."
            )
        question_ids.add(question_id)
        questions.append({
            "id": question_id,
            "category": str(category or "").strip(),
            "question": str(question_text).strip(),
            "choices": choices,
            "correct_answer": correct_answer,
        })
    return questions


def _test_identifier(value):
    return re.sub(r"[^a-z0-9]+", "-", str(value).strip().lower()).strip("-") or "test"


def load_workbook_tests():
    workbook_path = app.config.get(
        "FINAL_ENGLISH_TEST_WORKBOOK",
        os.path.join(app.root_path, "static", "data", "tests.xlsx"),
    )
    try:
        workbook = load_workbook(workbook_path, data_only=True)
        try:
            tests = []
            used_ids = set()

            def add_test(
                source_name,
                display_name,
                rows,
                header=None,
                allow_legacy_layout=True,
            ):
                questions = _parse_test_questions(
                    rows,
                    source_name,
                    header,
                    allow_legacy_layout,
                )
                if not questions:
                    return
                identifier = _test_identifier(source_name)
                base_identifier = identifier
                suffix = 2
                while identifier in used_ids:
                    identifier = f"{base_identifier}-{suffix}"
                    suffix += 1
                used_ids.add(identifier)
                tests.append({
                    "id": identifier,
                    "name": display_name,
                    "questions": questions,
                })

            for worksheet in workbook.worksheets:
                if worksheet.tables:
                    for table in worksheet.tables.values():
                        table_rows = [
                            [cell.value for cell in row]
                            for row in worksheet[table.ref]
                        ]
                        if table_rows:
                            add_test(
                                f"{worksheet.title}-{table.displayName}",
                                f"{worksheet.title} — {table.displayName}",
                                table_rows[1:],
                                table_rows[0],
                                allow_legacy_layout=False,
                            )
                    continue

                worksheet_rows = list(
                    worksheet.iter_rows(values_only=True)
                )
                if not worksheet_rows:
                    continue
                header = worksheet_rows[0]
                display_name = (
                    "Final English Test"
                    if worksheet.title == "English Test"
                    else worksheet.title
                )
                source_name = (
                    "final-english"
                    if worksheet.title == "English Test"
                    else worksheet.title
                )
                questions = _parse_test_questions(
                    worksheet_rows[1:],
                    worksheet.title,
                    header,
                )
                if questions:
                    identifier = _test_identifier(source_name)
                    base_identifier = identifier
                    suffix = 2
                    while identifier in used_ids:
                        identifier = f"{base_identifier}-{suffix}"
                        suffix += 1
                    used_ids.add(identifier)
                    tests.append({
                        "id": identifier,
                        "name": display_name,
                        "questions": questions,
                    })
            if not tests:
                raise ValueError("The workbook does not contain any test questions.")
            return tests
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException) as error:
        app.logger.exception("Could not load the test workbook.")
        raise RuntimeError("The tests are temporarily unavailable.") from error


def load_final_english_test():
    return _primary_workbook_test(load_workbook_tests())["questions"]


def _primary_workbook_test(tests):
    return next(
        (
            test
            for test in tests
            if test["id"] == "final-english"
            or test["name"] == "Final English Test"
            or test["name"].startswith("English Test — ")
        ),
        tests[0],
    )


def _find_workbook_test(test_id):
    tests = load_workbook_tests()
    test = next((item for item in tests if item["id"] == test_id), None)
    if test is None and test_id == "final-english":
        return _primary_workbook_test(tests)
    return test


def valid_test_email(email):
    if not email or len(email) > 254 or email.count("@") != 1:
        return False
    local, domain = email.rsplit("@", 1)
    if (
        not local
        or len(local) > 64
        or local.startswith(".")
        or local.endswith(".")
        or ".." in local
        or not re.fullmatch(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+", local)
    ):
        return False
    labels = domain.split(".")
    return (
        len(labels) >= 2
        and all(
            len(label) <= 63
            and re.fullmatch(r"[A-Za-z0-9](?:[A-Za-z0-9-]*[A-Za-z0-9])?", label)
            for label in labels
        )
    )


def is_logged_in():
    return "username" in session and "role" in session


def require_login():
    if not is_logged_in():
        return redirect(url_for("login"))
    return None


def local_next_url(value):
    if not value.startswith("/") or value.startswith("//"):
        return url_for("home")
    return value


def current_student():
    user_id = session.get("user_id")
    if not user_id:
        return None
    with closing(connect_student_database()) as connection, connection:
        return connection.execute(
            "SELECT id, name, last_name, email, role FROM student_users WHERE id = ?",
            (user_id,),
        ).fetchone()

# ============================================================
# JSON DATA
# ============================================================

def load_data(file_path):

    import os

    full_path = os.path.join(
        app.root_path,
        file_path
    )

    try:

        with open(
            full_path,
            'r',
            encoding='utf-8'
        ) as f:

            return json.load(f)

    except FileNotFoundError:

        return {}

# ============================================================
# LOGIN
# ============================================================

@app.route("/login", methods=["GET", "POST"])
def login():
    message = ""
    next_url = request.args.get("next", "")
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        next_url = request.form.get("next", next_url)

        user = USERS.get(username)
        registered_user = None
        if not user:
            with closing(connect_student_database()) as connection, connection:
                registered_user = connection.execute(
                    "SELECT * FROM student_users WHERE email = ? OR username = ?",
                    (username.strip().lower(), username.strip().lower())
                ).fetchone()

        configured_password_matches = user and user["password"] == password
        registered_password_matches = (
            registered_user
            and registered_user["password_hash"]
            and check_password_hash(registered_user["password_hash"], password)
        )
        if configured_password_matches or registered_password_matches:
            session.clear()
            session["username"] = username
            session["role"] = (
                user["role"]
                if user
                else "enrolled student"
                if registered_user["role"] == "student"
                else registered_user["role"]
            )
            if registered_user:
                session["user_id"] = registered_user["id"]
                session["display_name"] = f"{registered_user['name']} {registered_user['last_name']}"
            else:
                session.pop("user_id", None)
                session["display_name"] = username
            return redirect(local_next_url(next_url))
        else:
            message = "Invalid email or password."

    return render_template("login.html", message=message, next_url=next_url)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/register", methods=["GET", "POST"])
def register():
    message = ""
    next_url = request.args.get("next", "")
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        last_name = request.form.get("last_name", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        next_url = request.form.get("next", next_url)
        if not all([name, last_name, email, password]):
            message = "All fields are required."
        elif len(name) > 120 or len(last_name) > 120:
            message = "Names must be 120 characters or fewer."
        elif not valid_test_email(email):
            message = "Enter a valid email address."
        elif len(password) < 8:
            message = "Use a password with at least 8 characters."
        else:
            with closing(connect_student_database()) as connection, connection:
                try:
                    connection.execute(
                        "INSERT INTO student_users (name, last_name, email, password_hash, role) VALUES (?, ?, ?, ?, ?)",
                        (name, last_name, email, generate_password_hash(password), "student")
                    )
                    connection.commit()
                except database_integrity_errors():
                    message = "That email is already registered."
                else:
                    session.clear()
                    session["username"] = email
                    session["role"] = "enrolled student"
                    session["user_id"] = connection.execute(
                        "SELECT id FROM student_users WHERE email = ?", (email,)
                    ).fetchone()["id"]
                    session["display_name"] = f"{name} {last_name}"
                    if next_url:
                        return redirect(local_next_url(next_url))
                    return redirect(url_for("verb_conjugation_game"))
    return render_template("register.html", message=message, next_url=next_url)


@app.route("/verb-conjugation-game")
@app.route("/verb_conjugation_game")
def verb_conjugation_game():
    guest_mode = request.args.get("guest") == "1"
    saved_user = bool(session.get("user_id")) and not guest_mode
    return render_template(
        "verb_conjugation_game.html",
        display_name=session.get("display_name", "Guest") if saved_user else "Guest",
        saved_user=saved_user,
        guest_mode=guest_mode,
        admin_mode=session.get("role") == "admin"
    )


@app.route("/api/conjugation-progress", methods=["GET", "POST"])
def conjugation_progress():
    user_id = session.get("user_id")
    if not user_id:
        return {"saved": False, "progress": {}}

    if request.method == "POST":
        payload = request.get_json(silent=True) or {}
        group_name = str(payload.get("group_name", "")).strip()
        completed_series = max(0, int(payload.get("completed_series", 0)))
        current_verb = max(0, int(payload.get("current_verb", 0)))
        score = max(0, int(payload.get("score", 0)))
        if not group_name:
            return {"error": "group_name is required"}, 400
        with closing(connect_student_database()) as connection, connection:
            connection.execute("""
                INSERT INTO conjugation_progress (user_id, group_name, completed_series, current_verb, score)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(user_id, group_name) DO UPDATE SET
                    completed_series = excluded.completed_series,
                    current_verb = excluded.current_verb,
                    score = excluded.score
            """, (user_id, group_name, completed_series, current_verb, score))
            connection.commit()
        return {"saved": True}

    with closing(connect_student_database()) as connection, connection:
        rows = connection.execute("""
            SELECT group_name, completed_series, current_verb, score
            FROM conjugation_progress
            WHERE user_id = ?
        """, (user_id,)).fetchall()
    return {
        "saved": True,
        "progress": {
            row["group_name"]: {
                "completedSeries": row["completed_series"],
                "currentVerb": row["current_verb"],
                "score": row["score"],
            }
            for row in rows
        }
    }


@app.route("/api/conjugation-data")
def conjugation_data():
    workbook_path = os.path.join(
        app.root_path, "static", "data", "verbs_master_table.xlsx"
    )
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Verbs"]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [header for header in header_row if header]
        required_headers = {"parent_category", "series_id", "series_title"}
        if not required_headers.issubset(headers):
            return {"error": "The verbs master table has an invalid header row."}, 500

        groups = {}
        series_by_key = {}
        metadata_headers = {"parent_category", "series_id", "series_title"}
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if not row or all(value is None for value in row):
                continue
            values = {
                header: row[index] if index < len(row) else None
                for index, header in enumerate(header_row)
                if header
            }
            category = values.get("parent_category")
            if not category:
                continue

            group = groups.setdefault(category, {"series": []})
            series_key = (category, values.get("series_id"))
            series = series_by_key.get(series_key)
            if series is None:
                series = {
                    "id": values.get("series_id") or "",
                    "title": values.get("series_title") or "",
                    "verbs": [],
                }
                series_by_key[series_key] = series
                group["series"].append(series)

            series["verbs"].append({
                header: "" if value is None else value
                for header, value in values.items()
                if header not in metadata_headers
            })
        return app.response_class(
            json.dumps(groups, ensure_ascii=False), mimetype="application/json"
        )
    finally:
        workbook.close()



# ============================================================
# SUPPORT DATABASE
# ============================================================

def init_support_database():

    conn = sqlite3.connect("support_messages.db")

    conn.execute("""
        CREATE TABLE IF NOT EXISTS support_messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            email TEXT NOT NULL,
            subject TEXT NOT NULL,
            message TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    conn.commit()
    conn.close()


init_support_database()


# ============================================================
# HOME
# ============================================================

@app.route("/")
def home():
    return render_template("index.html")



# ============================================================
# VOCABULARY SETS
# ============================================================

VOCABULARY_SETS = {

    "basic-adjectives": {
        "title": "Basic Adjectives",
        "description": "Common adjectives for describing people, things, and situations.",
        "words": [
            "good = bueno",
            "bad = malo",
            "easy = fácil",
            "difficult = difícil",
            "big = grande",
            "small = pequeño",
            "tall = alto",
            "short = corto",
            "high = alto",
            "low = bajo",
            "hot = caliente",
            "cold = frío",
            "fast = rápido",
            "slow = lento",
            "happy = feliz",
            "sad = triste",
        ]
    },

    "subject-object-clauses": {
        "title": "Sub and Obj Clause Vocab",
        "description": "Words used to begin subject and object clauses.",
        "words": [
            "who = quien",
            "whose = de quien sea",
            "what = que",
            "when = cuando",
            "where = donde",
            "why = por qué",
            "how = como",
            "that = que", 
        ]
    },

  "subject-object-clauses else": {
        "title": "Sub and Obj Clause Vocab (else)",
        "description": "Words used to begin subject and object clauses with (else).",
        "words": [
            "who else = quien mas",
            "what else = que mas",
            "when else = cuando mas",
            "where else = donde mas",
            "why else = por qué mas / por que otra razón",
            "how else = como mas / de que otra manera"
        ]
    },

    "subject-object-clauses ever": {
        "title": "Sub and Obj Clause Vocab (ever)",
        "description": "Words used to begin subject and object clauses with (ever).",
        "words": [
            "whoever = quien sea",
            "whatever = que sea",
            "whenever = cuando sea",
            "wherever = donde sea",
            "whyever = por la razon qué sea",
            "however = como sea"
        ]
    },

    "everyday-vocabulary": {
        "title": "Everyday Vocabulary",
        "description": "Useful words that appear constantly in everyday English.",
        "words": [
            "thing = cosa",
            "way = camino, manera, forma",
            "place = lugar",
            "time = tiempo",
            "people = gente",
            "stuff = cosas"
        ]
    },

    "everyday-english": {
        "title": "Everyday English",
        "description": "Useful English vocabulary for everyday situations.",
        "words": [
            "lucky = afortunado",
            "unusual = no común",
            "silly = tonto",
            "strong = fuerte",
            "easy = fácil",
            "difficult = difícil",
            "busy = ocupado",
            "quiet = tranquilo",
            "careful = cuidadoso",
            "important = importante"
        ]
    },

    "essential-verbs": {
        "title": "Essential Verbs",
        "description": "High-frequency verbs used to build everyday English sentences.",
        "words": [
            "be = ser / estar",
            "have = tener",
            "do = hacer",
            "go = ir",
            "make = hacer / crear",
            "take = tomar / llevar",
            "come = venir",
            "put = poner",
            "leave = irse / dejar",
            "run = correr",
            "say = decir",
            "tell = contar / decir",
            "call = llamar"
        ]
    },

    "work": {
        "title": "Work Vocabulary",
        "description": "Practical vocabulary for talking about work, jobs, and daily responsibilities.",
        "words": [
            "job = trabajo",
            "shift = turno",
            "worker = trabajador",
            "manager = gerente",
            "meeting = reunión",
            "schedule = horario"
        ]
    },

    "colors": {
        "title": "Colors",
        "description": "Common colors for describing people, objects, clothes, and places.",
        "words": [
            "red = rojo",
            "blue = azul",
            "green = verde",
            "yellow = amarillo",
            "orange = naranja",
            "purple = morado",
            "pink = rosa",
            "brown = café",
            "black = negro",
            "white = blanco",
            "gray = gris",
            "gold = dorado",
            "silver = plateado"
        ]
    },

    "numbers": {
        "title": "Numbers",
        "description": "Numbers from 1 to 10",
        "words": [
            "1 = one wan",
            "2 = two tu",
            "3 = three thri",
            "4 = four for",
            "5 = five faiv",
            "6 = six siks",
            "7 = seven",
            "8 = eight eit",
            "9 = nine nain",
            "10 = ten ten"
        ]
    },

}


# ============================================================
# VOCABULARY PAGE
# ============================================================

@app.route('/vocab_drill')
def vocab():

    access_error = require_login()
    if access_error:
        return access_error

    return render_template(
        'vocab_drill.html',
        vocabulary_sets=VOCABULARY_SETS
    )


# ============================================================
# LESSONS
# ============================================================

@app.route('/lessons')
def lessons():

    access_error = require_login()
    if access_error:
        return access_error

    import os

    lessons_folder = os.path.join(
        app.root_path,
        'lessons'
    )

    lesson_files = []

    if os.path.exists(lessons_folder):

        for filename in os.listdir(lessons_folder):

            if filename.lower().endswith('.pdf'):

                lesson_files.append(filename)

    lesson_files.sort()

    return render_template(
        'lessons.html',
        lessons=lesson_files
    )


# ============================================================
# OPEN LESSON PDF
# ============================================================

@app.route('/lesson-pdf/<path:filename>')
def lesson_pdf(filename):

    import os

    lessons_folder = os.path.join(
        app.root_path,
        'lessons'
    )

    return send_from_directory(
        lessons_folder,
        filename
    )


# ============================================================
# COURSE
# ============================================================

@app.route('/course')
def course():

    access_error = require_login()
    if access_error:
        return access_error

    return render_template('course.html')


@app.route("/test-center")
def test_center():
    try:
        tests = load_workbook_tests()
    except RuntimeError:
        abort(503, description="The tests are temporarily unavailable.")
    return render_template("test_center.html", tests=tests)


@app.route("/tests/final-english/register", methods=["GET", "POST"])
def final_english_register():
    try:
        test = _primary_workbook_test(load_workbook_tests())
    except RuntimeError:
        abort(503, description="The tests are temporarily unavailable.")
    return redirect(url_for("final_english_take", test_id=test["id"]))


@app.route("/tests/<test_id>/take")
def final_english_take(test_id):
    try:
        test = _find_workbook_test(test_id)
    except RuntimeError:
        abort(503, description="The tests are temporarily unavailable.")
    if test is None:
        abort(404)
    questions = test["questions"]
    public_questions = [
        {
            "id": question["id"],
            "category": question["category"],
            "question": question["question"],
            "choices": question["choices"],
        }
        for question in questions
    ]
    return render_template(
        "test_take.html",
        questions=public_questions,
        test_id=test["id"],
        test_name=test["name"],
    )


@app.route("/tests/<test_id>/submit", methods=["POST"])
def final_english_submit(test_id):
    if not request.is_json:
        return {"error": "A JSON answer report is required."}, 400
    submission = request.get_json(silent=True)
    if not isinstance(submission, dict):
        return {"error": "The submitted answers are invalid."}, 400
    if (
        isinstance(submission.get("answers"), dict)
        and set(submission).issubset({"answers", "name"})
    ):
        answers = submission.get("answers")
        participant_name = submission.get("name", "")
    else:
        answers = submission
        participant_name = ""
    if not isinstance(participant_name, str) or len(participant_name.strip()) > 120:
        return {"error": "Names must be 120 characters or fewer."}, 400
    if not isinstance(answers, dict):
        return {"error": "The submitted answers are invalid."}, 400
    try:
        test = _find_workbook_test(test_id)
    except RuntimeError:
        abort(503, description="The tests are temporarily unavailable.")
    if test is None:
        abort(404)
    questions = test["questions"]

    expected_ids = {question["id"] for question in questions}
    if set(answers) != expected_ids or any(
        not isinstance(answer, str)
        or answer not in question["choices"]
        for question in questions
        for answer in [answers.get(question["id"])]
    ):
        return {"error": "Answer every question before submitting."}, 400

    report = []
    score = 0
    for question in questions:
        selected_answer = answers[question["id"]]
        is_correct = selected_answer == question["correct_answer"]
        score += int(is_correct)
        report.append({
            "question_id": question["id"],
            "category": question["category"],
            "question": question["question"],
            "choices": question["choices"],
            "selected_answer": selected_answer,
            "correct_answer": question["correct_answer"],
            "is_correct": is_correct,
        })

    incorrect_questions = [
        {
            "question_number": index,
            "question": result["question"],
            "selected_answer": result["choices"][result["selected_answer"]],
        }
        for index, result in enumerate(report, start=1)
        if not result["is_correct"]
    ]

    student = current_student()
    name = participant_name.strip()
    if not name and student:
        name = f"{student['name']} {student['last_name']}".strip()
    connection = connect_test_results_database()
    try:
        connection.execute(
            """
            INSERT INTO test_results (
                user_id, test_name, name, email, score, question_count, report_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                student["id"] if student else None,
                test["name"],
                name or "Guest",
                student["email"] if student else "",
                score,
                len(questions),
                json.dumps(report, ensure_ascii=False),
            ),
        )
        connection.commit()
    finally:
        connection.close()
    return {
        "submitted": True,
        "score": score,
        "question_count": len(questions),
        "incorrect_questions": incorrect_questions,
    }


def load_admin_students():
    with closing(connect_student_database()) as connection, connection:
        users = [
            dict(row)
            for row in connection.execute(
                "SELECT id, name, last_name, email FROM student_users ORDER BY name, last_name, id"
            )
        ]
        approvals = connection.execute(
            """
            SELECT user_id, category, approved
            FROM student_test_approvals
            """
        ).fetchall()
        conjugation_progress = connection.execute(
            """
            SELECT user_id, group_name, completed_series, score
            FROM conjugation_progress ORDER BY group_name
            """
        ).fetchall()
        game_tables = {
            row["name"] for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        phantom_progress = []
        if "phantom_pronouns_progress" in game_tables:
            phantom_progress = connection.execute(
                """
                SELECT user_id, level_key, best_score, completed
                FROM phantom_pronouns_progress ORDER BY level_key
                """
            ).fetchall()

    by_id = {user["id"]: user for user in users}
    profile_user_ids = {f"user:{user['id']}": user["id"] for user in users}
    for user in users:
        user["assessment_approved"] = False
        user["game_progress"] = []
    for approval in approvals:
        if approval["category"] == "assessment" and approval["user_id"] in by_id:
            by_id[approval["user_id"]]["assessment_approved"] = bool(approval["approved"])
    for progress in conjugation_progress:
        if progress["user_id"] in by_id:
            by_id[progress["user_id"]]["game_progress"].append({
                "game": "Verb conjugation",
                "level": progress["group_name"],
                "score": progress["score"],
                "completed": f"{progress['completed_series']} series completed",
            })
    for progress in phantom_progress:
        if progress["user_id"] in by_id:
            by_id[progress["user_id"]]["game_progress"].append({
                "game": "Phantom Pronouns",
                "level": progress["level_key"],
                "score": progress["best_score"],
                "completed": "Completed" if progress["completed"] else "In progress",
            })

    connection = connect_word_order_database()
    try:
        progress_rows = connection.execute(
            """
            SELECT p.profile_key, p.level_id, p.best_score, p.completed_at,
                l.level_name_en
            FROM word_order_progress AS p
            LEFT JOIN levels AS l ON l.level_id = p.level_id
            WHERE p.profile_key LIKE 'user:%'
            ORDER BY p.level_id
            """
        ).fetchall()
    finally:
        connection.close()
    for progress in progress_rows:
        user_id = profile_user_ids.get(progress["profile_key"])
        if user_id is not None:
            by_id[user_id]["game_progress"].append({
                "game": "Word Order",
                "level": progress["level_name_en"] or f"Level {progress['level_id']}",
                "score": progress["best_score"],
                "completed": "Completed" if progress["completed_at"] else "In progress",
            })
    return users


@app.route("/admin/test-results", methods=["GET", "POST"])
def admin_test_results():
    if request.method == "POST":
        password = request.form.get("password", "")
        if check_password_hash(TEST_RESULTS_ADMIN_PASSWORD_HASH, password):
            session["test_results_admin_authenticated"] = True
            return redirect(url_for("admin_test_results"))
        return render_template(
            "test_results_admin.html",
            authenticated=False,
            message="Incorrect password.",
        ), 401

    if not session.get("test_results_admin_authenticated"):
        return render_template(
            "test_results_admin.html",
            authenticated=False,
            message="",
        )

    connection = connect_test_results_database()
    try:
        results = connection.execute(
            "SELECT * FROM test_results ORDER BY created_at DESC, id DESC"
        ).fetchall()
    finally:
        connection.close()
    response = app.make_response(render_template(
        "test_results_admin.html",
        authenticated=True,
        students=load_admin_students(),
        results=[
            {**dict(result), "report": json.loads(result["report_json"])}
            for result in results
        ],
    ))
    response.headers["Cache-Control"] = "no-store"
    return response


@app.route("/admin/test-results/users/<int:user_id>/approvals", methods=["POST"])
def admin_update_test_approval(user_id):
    if not session.get("test_results_admin_authenticated"):
        abort(403)
    category = request.form.get("category", "")
    approved_value = request.form.get("approved", "")
    if category != "assessment" or approved_value not in {"0", "1"}:
        abort(400, description="Choose a valid test category and approval state.")
    with closing(connect_student_database()) as connection, connection:
        user_exists = connection.execute(
            "SELECT 1 FROM student_users WHERE id = ?",
            (user_id,),
        ).fetchone()
        if not user_exists:
            abort(404)
        connection.execute(
            """
            INSERT INTO student_test_approvals (user_id, category, approved)
            VALUES (?, ?, ?)
            ON CONFLICT(user_id, category) DO UPDATE SET
                approved = excluded.approved,
                updated_at = CURRENT_TIMESTAMP
            """,
            (user_id, category, int(approved_value)),
        )
        connection.commit()
    return redirect(url_for("admin_test_results"))


@app.route("/admin/test-results/logout", methods=["POST"])
def admin_test_results_logout():
    session.pop("test_results_admin_authenticated", None)
    return redirect(url_for("admin_test_results"))


# ============================================================
# MODULES
# ============================================================

@app.route('/module/<module_id>')
def interactive_module(module_id):

    access_error = require_login()
    if access_error:
        return access_error

    # --------------------------------------------------------
    # LOAD JSON DATA
    # --------------------------------------------------------

    videos_data = load_data('static/data/videos.json')
    quizzes_data = load_data('static/data/quizzes.json')


    # --------------------------------------------------------
    # FIND THE REQUESTED MODULE
    # --------------------------------------------------------

    video_module = videos_data.get(module_id)
    quiz_module = quizzes_data.get(module_id)

    if not video_module and not quiz_module:
        return f"URL module_id = [{module_id}]<br>JSON keys = {list(videos_data.keys())}"

    content_module = video_module or quiz_module


    # --------------------------------------------------------
    # CREATE THE MODULE
    # --------------------------------------------------------

    module = {

        "title": content_module.get(
            "title",
            module_id
        ),

        "description": content_module.get(
            "description",
            ""
        ),

        "series": []

    }


    # --------------------------------------------------------
    # INDEX THE QUIZ SERIES
    # --------------------------------------------------------

    quiz_series = {}

    if quiz_module:

        for series in quiz_module.get(
            "series",
            []
        ):

            quiz_series[
                series.get("id")
            ] = series


    # --------------------------------------------------------
    # COMBINE VIDEO + QUIZ SERIES
    # --------------------------------------------------------

    video_series = {
        series.get("id"): series
        for series in (video_module or {}).get("series", [])
    }

    all_series_ids = list(
        dict.fromkeys(
            list(video_series.keys()) +
            list(quiz_series.keys())
        )
    )

    for series_id in all_series_ids:

        video_series_data = video_series.get(series_id, {})
        quiz_series_data = quiz_series.get(series_id, {})


        combined_series = {

            "id": series_id,

            "title": video_series_data.get(
                "title",
                quiz_series_data.get("title", f"Series {series_id}")
            ),

            "description": video_series_data.get(
                "description",
                quiz_series_data.get("description", "")
            ),

            "videos": video_series_data.get(
                "videos",
                []
            ),

            "quiz": quiz_series_data.get(
                "questions",
                []
            ),

            "audio": video_series_data.get(
                "audio",
                []
            ),

            "images": video_series_data.get(
                "images",
                []
            )

        }


        module["series"].append(
            combined_series
        )


    # --------------------------------------------------------
    # SEND COMPLETE MODULE TO HTML
    # --------------------------------------------------------

    return render_template(
        'module.html',
        module=module
    )
# ============================================================
# GAMES
# ============================================================

@app.route('/games')
def games_dashboard():

    return render_template('games.html')

    if __name__ == "__main__":
        app.run(host='0.0.0.0.', por=5000, debug=Ture)

# ============================================================
# ABOUT
# ============================================================

@app.route('/about')
def about():
    return render_template('about.html')


# ============================================================
# LIVE CLASSES
# ============================================================

@app.route('/live_classes')
def live_classes():
    return render_template('live_classes.html')


# ============================================================
# CONTACT
# ============================================================

@app.route('/contact')
def contact():
    return render_template('contact.html')

@app.route("/tech-support", methods=["GET", "POST"])
def tech_support():

    if request.method == "POST":

        name = request.form["name"]
        email = request.form["email"]
        subject = request.form["subject"]
        message = request.form["message"]

        conn = sqlite3.connect("support_messages.db")

        conn.execute("""
            INSERT INTO support_messages
            (name, email, subject, message)
            VALUES (?, ?, ?, ?)
        """, (name, email, subject, message))

        conn.commit()
        conn.close()

        return render_template(
            "tech_support.html",
            success=True
        )

    return render_template("tech_support.html")

@app.route("/support_messages")
def support_messages():

    access_error = require_login()
    if access_error:
        return access_error

    if session["role"] == "enrolled student":
        return "Access denied"

    conn = sqlite3.connect("support_messages.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id, name, email, subject, message, created_at
        FROM support_messages
        ORDER BY created_at DESC
    """)

    messages = cursor.fetchall()
    conn.close()

    return render_template(
        "support_messages.html",
        messages=messages
    )

# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG", "0") == "1")
