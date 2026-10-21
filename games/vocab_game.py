from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path
import re
from urllib.parse import quote, unquote, urlsplit, urlunsplit

from flask import Blueprint, jsonify, redirect, render_template, request, session, url_for
from openpyxl import load_workbook

from gallery import exercise_image_url, public_audio_url


vocab_bp = Blueprint("vocab_game", __name__, url_prefix="/vocab")

REPOSITORY_DIR = Path(__file__).resolve().parents[1]
STATIC_DIR = REPOSITORY_DIR / "static"
VOCAB_IMAGE_DIR = STATIC_DIR / "images" / "vocab_game"
VOCAB_AUDIO_DIR = STATIC_DIR / "audio" / "vocab_game"
VOCAB_WORKBOOK = STATIC_DIR / "data" / "Vocab_Game.xlsx"
AUDIO_EXTENSIONS = {".m4a", ".mp3", ".ogg", ".opus", ".wav"}
AUDIO_EXTENSION_PRIORITY = {
    ".mp3": 0,
    ".opus": 1,
    ".ogg": 2,
    ".wav": 3,
    ".m4a": 4,
}
_EXPLICIT_LABEL_ALIASES = {"sing": {"sink"}}


def _normalized_label(value):
    return " ".join(re.findall(r"[a-z0-9]+", str(value).casefold()))


def _label_aliases(value):
    aliases = []
    for part in re.split(r"\s*/\s*|\s+or\s+", str(value), flags=re.IGNORECASE):
        normalized = _normalized_label(part)
        if normalized and normalized not in aliases:
            aliases.append(normalized)
    return aliases


def _labels_match(left, right):
    left_aliases = _label_aliases(left)
    right_aliases = _label_aliases(right)
    for left_alias in left_aliases:
        for right_alias in right_aliases:
            if left_alias == right_alias:
                return True
            if right_alias in _EXPLICIT_LABEL_ALIASES.get(left_alias, set()):
                return True
            if left_alias in _EXPLICIT_LABEL_ALIASES.get(right_alias, set()):
                return True
            if min(len(left_alias), len(right_alias)) >= 4:
                if SequenceMatcher(None, left_alias, right_alias).ratio() >= 0.8:
                    return True
    return False


def _local_asset_path(asset_url):
    parsed_url = urlsplit(str(asset_url))
    if parsed_url.scheme or parsed_url.netloc:
        return None

    asset_path = unquote(parsed_url.path)
    if not asset_path.startswith("/static/"):
        return None

    local_path = (REPOSITORY_DIR / asset_path.lstrip("/")).resolve()
    try:
        local_path.relative_to(STATIC_DIR.resolve())
    except ValueError:
        return None

    return local_path if local_path.is_file() else None


def _asset_url_from_filename(filename, directory, expected_dir):
    if not filename:
        return None
    asset_path = (directory / Path(str(filename)).name).resolve()
    try:
        asset_path.relative_to(expected_dir.resolve())
    except ValueError:
        return None
    if not asset_path.is_file():
        return None
    relative_path = asset_path.relative_to(STATIC_DIR.resolve()).as_posix()
    return f"/static/{relative_path}"


def _reference_rows(sheet):
    references = []
    for label, asset_url, *_ in sheet.iter_rows(
        min_row=2, max_col=2, values_only=True
    ):
        if not label or not asset_url:
            continue
        parsed_url = urlsplit(str(asset_url).strip())
        if (
            parsed_url.scheme.casefold() not in {"http", "https"}
            or not parsed_url.netloc
        ):
            raise ValueError(
                f"Vocabulary workbook sheet '{sheet.title}' has an invalid media URL "
                f"for '{label}'."
            )
        path = re.sub(r"^/\s+images/", "/images/", parsed_url.path.strip())
        normalized_url = urlunsplit(
            (
                parsed_url.scheme,
                parsed_url.netloc,
                quote(path, safe="/%:@"),
                parsed_url.query,
                parsed_url.fragment,
            )
        )
        references.append(
            {
                "label": str(label).strip(),
                "url": normalized_url,
                "filename": Path(unquote(path)).name,
                "order": len(references),
            }
        )
    return references


def _reference_index(references, include_suffix=False):
    index = {}
    for reference in references:
        for alias in _label_aliases(reference["label"]):
            indexed_aliases = [alias]
            if include_suffix:
                parts = alias.split()
                indexed_aliases.extend(
                    " ".join(parts[start:]) for start in range(1, len(parts))
                )
            for indexed_alias in indexed_aliases:
                index.setdefault(indexed_alias, []).append(reference)
    return index


def _best_audio_reference(word, audio_references, audio_index):
    exact_matches = {
        reference["order"]: reference
        for alias in _label_aliases(word)
        for reference in audio_index.get(alias, [])
    }
    if exact_matches:
        normalized_word = _normalized_label(word)
        return min(
            exact_matches.values(),
            key=lambda reference: (
                _normalized_label(reference["label"]) != normalized_word,
                len(_normalized_label(reference["label"])),
            ),
        )
    return next(
        (
            reference
            for reference in audio_references
            if _labels_match(reference["label"], word)
        ),
        None,
    )


def _matching_image_references(word, image_references, image_index):
    matches = {
        reference["order"]: reference
        for alias in _label_aliases(word)
        for reference in image_index.get(alias, [])
    }
    if not matches:
        matches = {
            reference["order"]: reference
            for reference in image_references
            if _labels_match(reference["label"], word)
        }
    return sorted(matches.values(), key=lambda reference: reference["order"])


def _media_row_assets(row, image_paths):
    audio_url = row[6]
    if audio_url and _local_asset_path(audio_url) is None:
        if not urlsplit(str(audio_url)).scheme:
            audio_url = None
    if not audio_url:
        audio_url = _asset_url_from_filename(row[5], VOCAB_AUDIO_DIR, VOCAB_AUDIO_DIR)
    if not audio_url:
        return None

    image_urls = []
    image_url = row[8]
    if image_url and (
        _local_asset_path(image_url) is not None or urlsplit(str(image_url)).scheme
    ):
        image_urls.append(str(image_url))
    elif row[7]:
        linked_image = _asset_url_from_filename(row[7], VOCAB_IMAGE_DIR, VOCAB_IMAGE_DIR)
        if linked_image:
            image_urls.append(linked_image)

    labels = []
    for value in (row[1], row[2], Path(str(row[5] or "")).stem):
        if value:
            labels.extend(_label_aliases(value))
    labels = list(dict.fromkeys(labels))

    for image_path in image_paths:
        image_label = _normalized_label(image_path.stem)
        if any(
            image_label == label or image_label.endswith(f" {label}")
            for label in labels
        ):
            relative_path = image_path.relative_to(STATIC_DIR).as_posix()
            image_urls.append(f"/static/{relative_path}")

    image_urls = list(dict.fromkeys(image_urls))
    if not image_urls:
        return None

    try:
        audio_path = _local_asset_path(audio_url)
    except (OSError, ValueError):
        audio_path = None
    if audio_path is not None and audio_path.suffix.casefold() not in AUDIO_EXTENSIONS:
        return None

    return {
        "audio": str(audio_url),
        "images": image_urls,
        "labels": labels,
        "row": row,
    }


def _workbook_rows():
    if not VOCAB_WORKBOOK.is_file():
        raise FileNotFoundError(f"Vocabulary workbook not found: {VOCAB_WORKBOOK}")

    workbook = load_workbook(VOCAB_WORKBOOK, read_only=True, data_only=True)
    try:
        required_sheets = {"sets", "Audio Reference", "Video Reference"}
        if not required_sheets.issubset(workbook.sheetnames):
            raise ValueError(
                "Vocabulary workbook must contain 'sets', 'Audio Reference', "
                "and 'Video Reference' sheets."
            )

        set_rows = list(workbook["sets"].iter_rows(values_only=True))
        audio_references = _reference_rows(workbook["Audio Reference"])
        image_references = _reference_rows(workbook["Video Reference"])
    finally:
        workbook.close()

    if not set_rows:
        raise ValueError("Vocabulary workbook 'sets' sheet is empty.")
    if not audio_references or not image_references:
        raise ValueError("Vocabulary workbook media reference sheets are empty.")
    audio_index = _reference_index(audio_references)
    image_index = _reference_index(image_references, include_suffix=True)

    sets = {}
    table_ids = {}
    section_headers = []
    for index, row in enumerate(set_rows):
        if not row:
            continue
        section_match = re.fullmatch(r"set\s+(\d+)", str(row[0] or "").strip(), re.I)
        if section_match:
            section_headers.append((index, section_match.group(1)))
    if not section_headers:
        raise ValueError("Vocabulary workbook 'sets' sheet has no numbered sections.")

    for section_index, (header_index, table_id) in enumerate(section_headers):
        header = set_rows[header_index]
        next_header_index = (
            section_headers[section_index + 1][0]
            if section_index + 1 < len(section_headers)
            else len(set_rows)
        )
        column_sets = {}
        for column_index, set_value in enumerate(header[1:], start=1):
            if set_value is None or not str(set_value).strip():
                continue
            set_id = str(set_value).strip()
            column_sets[column_index] = set_id
            sets.setdefault(set_id, [])
            table_ids.setdefault(set_id, table_id)

        for row in set_rows[header_index + 1 : next_header_index]:
            for column_index, set_id in column_sets.items():
                value = row[column_index] if column_index < len(row) else None
                if value is None or not str(value).strip():
                    continue
                word = str(value).strip()
                if re.fullmatch(r"\d+(?:\.\d+)?", word):
                    continue
                if not any(
                    _normalized_label(word) == _normalized_label(item)
                    for item in sets[set_id]
                ):
                    sets[set_id].append(word)

    media = []
    for set_id, words in sets.items():
        for word in words:
            audio = _best_audio_reference(word, audio_references, audio_index)
            if not audio:
                continue
            matching_images = _matching_image_references(
                word, image_references, image_index
            )
            if not matching_images:
                continue
            for image in matching_images:
                row = (
                    set_id,
                    word,
                    None,
                    None,
                    None,
                    audio["filename"],
                    audio["url"],
                    image["filename"],
                    image["url"],
                    None,
                )
                assets = _media_row_assets(row, [])
                if assets:
                    media.append(assets)

    return sets, table_ids, media


def _row_matches_word(row, word):
    return any(_labels_match(word, label) for label in row["labels"])


def _row_set_ids(row):
    raw_set_ids = str(row["row"][0] or "")
    ids = set()
    for raw_id in re.split(r"[;,]", raw_set_ids):
        set_id = raw_id.strip()
        if set_id:
            ids.add(set_id)
    return ids


def _build_word(set_id, word, media_rows, item_index):
    ordered_rows = sorted(
        media_rows,
        key=lambda row: (
            AUDIO_EXTENSION_PRIORITY.get(
                Path(urlsplit(row["audio"]).path).suffix.casefold(), 10
            ),
            Path(urlsplit(row["audio"]).path).name.casefold(),
        ),
    )
    audio = ordered_rows[0]["audio"]
    images = list(dict.fromkeys(image for row in ordered_rows for image in row["images"]))
    if not images:
        return None

    normalized_id = _normalized_label(word).replace(" ", "-")
    return {
        "id": f"{set_id}:{normalized_id}:{item_index}",
        "english": word,
        "image": images[0],
        "images": images,
        "audio": audio,
    }


@lru_cache(maxsize=1)
def vocabulary_sets():
    sets, table_ids, media = _workbook_rows()
    set_items = {set_id: [] for set_id in sets}
    assigned_rows = {set_id: set() for set_id in sets}

    for set_id, words in sets.items():
        for word in words:
            matching_rows = [
                (index, row)
                for index, row in enumerate(media)
                if set_id in _row_set_ids(row) and _row_matches_word(row, word)
            ]
            if not matching_rows:
                continue
            item = _build_word(
                set_id, word, [row for _, row in matching_rows], len(set_items[set_id])
            )
            if item:
                set_items[set_id].append(item)
                assigned_rows[set_id].update(index for index, _ in matching_rows)

    for index, row in enumerate(media):
        target_set_ids = _row_set_ids(row) & set_items.keys()
        for set_id in target_set_ids:
            if index in assigned_rows[set_id]:
                continue
            label = row["row"][1] or row["row"][2]
            if not label:
                continue
            word = str(label).strip()
            matching_existing = next(
                (
                    item
                    for item in set_items[set_id]
                    if _labels_match(item["english"], word)
                ),
                None,
            )
            if matching_existing:
                matching_existing["images"] = list(
                    dict.fromkeys(matching_existing["images"] + row["images"])
                )
                continue
            item = _build_word(set_id, word, [row], len(set_items[set_id]))
            if item:
                set_items[set_id].append(item)
                assigned_rows[set_id].add(index)

    return {
        set_id: {
            "id": set_id,
            "title": f"Set {set_id}",
            "table_id": table_ids[set_id],
            "words": words,
        }
        for set_id, words in set_items.items()
    }


def thumbnail_words(words):
    formatted_words = []
    for word in words:
        formatted_word = {**word}
        formatted_word["images"] = []
        for image in word.get("images", [word["image"]]):
            parsed_image = urlsplit(image)
            image_prefix = "/images/vocab_game/"
            if parsed_image.scheme and image_prefix in parsed_image.path:
                image_path = parsed_image.path.split(image_prefix, 1)[1]
                local_image_url = (
                    "/static/images/vocab_game/"
                    + quote(unquote(image_path), safe="/")
                )
                thumbnail_url = exercise_image_url(local_image_url)
                formatted_word["images"].append(
                    thumbnail_url
                    if thumbnail_url != local_image_url
                    or _local_asset_path(local_image_url)
                    else image
                )
            else:
                formatted_word["images"].append(exercise_image_url(image))
        formatted_word["image"] = formatted_word["images"][0]
        audio = word["audio"]
        parsed_audio = urlsplit(audio)
        audio_prefix = "/audio/vocab_game/"
        if parsed_audio.scheme and audio_prefix in parsed_audio.path:
            audio_path = parsed_audio.path.split(audio_prefix, 1)[1]
            local_audio_url = (
                "/static/audio/vocab_game/"
                + quote(unquote(audio_path), safe="/")
            )
            public_url = public_audio_url(local_audio_url)
            formatted_word["audio"] = (
                public_url
                if public_url != local_audio_url or _local_asset_path(local_audio_url)
                else audio
            )
        else:
            formatted_word["audio"] = public_audio_url(audio)
        formatted_words.append(formatted_word)
    return formatted_words


@vocab_bp.route("/game")
def game_page():
    set_groups = {}
    for set_id, details in vocabulary_sets().items():
        table_id = details["table_id"]
        set_groups.setdefault(table_id, []).append(
            {
                "id": set_id,
                "title": details["title"],
                "word_count": len(details["words"]),
            }
        )
    return render_template(
        "games/vocab_game.html",
        set_groups=set_groups,
        set_title="Vocabulary Practice",
    )


@vocab_bp.route("/picture-match")
@vocab_bp.route("/picture-match/along")
@vocab_bp.route("/picture-match/french-audio")
@vocab_bp.route("/along")
def legacy_vocabulary_game():
    return redirect(url_for("vocab_game.game_page"))


@vocab_bp.route("/api/get-vocab-sets", methods=["GET"])
def get_vocab_sets():
    sets_list = [
        {
            "id": set_id,
            "title": details["title"],
            "table_id": details["table_id"],
            "word_count": len(details["words"]),
        }
        for set_id, details in vocabulary_sets().items()
    ]
    return jsonify(sets_list)


@vocab_bp.route("/api/get-set-details", methods=["GET"])
def get_set_details():
    set_id = request.args.get("id", "").strip()
    selected_set = vocabulary_sets().get(set_id)
    if selected_set is None:
        return jsonify({"error": "Set not found"}), 404
    if not selected_set["words"]:
        return jsonify({"error": "This set does not have playable image/audio pairs yet."}), 422

    session["active_set_id"] = set_id
    session["current_round"] = 1
    session["word_index"] = 0
    return jsonify(thumbnail_words(selected_set["words"]))
