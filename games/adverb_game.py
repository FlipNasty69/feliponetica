import os
import random

from flask import Blueprint, current_app, jsonify, render_template, request
from openpyxl import load_workbook

adverb_game_bp = Blueprint(
    "adverb_game", __name__, url_prefix="/adverb-game"
)

WORKBOOK_PATH = os.path.abspath(os.path.join(
    os.path.dirname(__file__), "..", "static", "data", "adverbs.xlsx"
))
_game_data_cache = None


def _display_text(value):
    return str(value or "").replace("_", " ").strip()


def _display_spanish(value):
    return " / ".join(_display_text(part) for part in str(value or "").split("|"))


def _read_sheet(workbook, sheet_name, required_headers):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Adverb workbook is missing the {sheet_name} sheet.")

    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    if not set(required_headers).issubset(headers):
        raise ValueError(f"Adverb workbook has an invalid {sheet_name} sheet.")
    return [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in values)
    ]


def get_game_data():
    global _game_data_cache
    if _game_data_cache is not None:
        return _game_data_cache

    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        categories = _read_sheet(
            workbook,
            "Categories",
            {"category_id", "category_key", "category_en"},
        )
        adverbs = _read_sheet(
            workbook,
            "Adverbs",
            {"adverb_id", "category_id", "adverb"},
        )
        meanings = _read_sheet(
            workbook,
            "Meaning_Game",
            {
                "adverb_id", "prompt_type", "correct_answer",
                "option_1", "option_2", "option_3", "option_4",
            },
        )
        placements = _read_sheet(
            workbook,
            "Placement_Game",
            {
                "adverb_id", "sentence_with_blank", "candidate_sentence",
                "is_correct", "position_label", "spanish_sentence",
            },
        )
    finally:
        workbook.close()

    category_by_id = {item["category_id"]: item for item in categories}
    adverb_by_id = {item["adverb_id"]: item for item in adverbs}
    category_options = [
        {
            "key": item["category_key"],
            "label": _display_text(item["category_en"]).replace("Adverbs of ", ""),
        }
        for item in categories
    ]

    meaning_questions = []
    for item in meanings:
        adverb = adverb_by_id.get(item["adverb_id"])
        if not adverb or adverb["category_id"] not in category_by_id:
            raise ValueError("Adverb workbook contains an unknown adverb category.")
        choices = [
            _display_spanish(item[f"option_{number}"])
            for number in range(1, 5)
        ]
        meaning_questions.append({
            "prompt": f"What does “{_display_text(adverb['adverb'])}” mean?",
            "adverb": _display_text(adverb["adverb"]),
            "category": category_by_id[adverb["category_id"]]["category_key"],
            "choices": choices,
            "answer": _display_spanish(item["correct_answer"]),
        })

    placement_questions = []
    for item in placements:
        adverb = adverb_by_id.get(item["adverb_id"])
        if not adverb or adverb["category_id"] not in category_by_id:
            raise ValueError("Adverb workbook contains an unknown adverb category.")
        sentence = _display_text(item["candidate_sentence"])
        if sentence and sentence[-1] not in ".?!":
            sentence += "."
        blank_sentence = _display_text(item["sentence_with_blank"])
        blank_sentence = blank_sentence.replace("adverb slot", "_____")
        placement_questions.append({
            "prompt": "Is the adverb in the right place?",
            "sentence": sentence,
            "blank_sentence": blank_sentence,
            "adverb": _display_text(adverb["adverb"]),
            "category": category_by_id[adverb["category_id"]]["category_key"],
            "position": _display_text(item["position_label"]).replace("_", " "),
            "spanish": _display_text(item["spanish_sentence"]),
            "answer": item["is_correct"] is True
            or str(item["is_correct"]).strip().lower() == "true",
        })

    _game_data_cache = {
        "categories": category_options,
        "meaning": meaning_questions,
        "placement": placement_questions,
    }
    return _game_data_cache


@adverb_game_bp.route("/game")
def game_page():
    return render_template(
        "games/adverb_game.html",
        categories=get_game_data()["categories"],
    )


@adverb_game_bp.route("/api/round")
def get_round():
    mode = request.args.get("mode", "meaning")
    if mode not in {"meaning", "placement"}:
        return jsonify({"error": "Choose a valid game mode."}), 400

    category = request.args.get("category", "all")
    count_value = request.args.get("count", "10")
    try:
        count = int(count_value)
    except ValueError:
        return jsonify({"error": "Question count must be a number."}), 400
    if count < 1 or count > 25:
        return jsonify({"error": "Question count must be between 1 and 25."}), 400

    data = get_game_data()
    valid_categories = {item["key"] for item in data["categories"]}
    if category != "all" and category not in valid_categories:
        return jsonify({"error": "Choose a valid adverb category."}), 400

    questions = data[mode]
    if category != "all":
        questions = [question for question in questions if question["category"] == category]
    if not questions:
        return jsonify({"error": "There are no questions for this selection."}), 404

    selected = [
        dict(question, choices=question["choices"][:])
        if mode == "meaning" else dict(question)
        for question in random.sample(questions, min(count, len(questions)))
    ]
    if mode == "meaning":
        for question in selected:
            random.shuffle(question["choices"])
    return jsonify({"mode": mode, "questions": selected})
