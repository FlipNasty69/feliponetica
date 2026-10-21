import os
import random
import re

from flask import Blueprint, current_app, jsonify, render_template, request
from openpyxl import load_workbook

question_game_bp = Blueprint(
    "question_game", __name__, url_prefix="/question-game"
)

WORKBOOK_PATH = os.path.abspath(os.path.join(
    os.path.dirname(__file__), "..", "static", "data", "question_game.xlsx"
))
_game_data_cache = None


def _text(value):
    return str(value).strip() if value is not None else ""


def _slug(value):
    slug = re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")
    return slug or "exercise"


def _is_generic_header(value):
    return not value or value.isdigit()


def _looks_like_slot_header(row):
    header_terms = (
        "auxiliary", "bully", "subject", "verb", "adverb", "object",
        "modifier", "complement",
    )
    values = [_text(value).casefold() for value in row[1:]]
    return not _text(row[0] if row else "") and any(
        value and any(term in value for term in header_terms)
        for value in values
    )


def _slot_labels(header, secondary_header, column_count):
    labels = []
    for index in range(1, column_count):
        label = _text(header[index]) if index < len(header) else ""
        if _is_generic_header(label):
            secondary = (
                _text(secondary_header[index])
                if index < len(secondary_header)
                else ""
            )
            label = secondary if secondary and not secondary.isdigit() else ""
        labels.append(None if label.casefold().startswith("note") else label or f"Word {index}")
    return labels


def _load_sheet_exercise(sheet, exercise_id):
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return None

    header = rows[0]
    secondary_header = (
        rows[1] if len(rows) > 1 and _looks_like_slot_header(rows[1]) else ()
    )
    if not _text(header[0] if header else ""):
        return None

    labels = _slot_labels(header, secondary_header, sheet.max_column)
    questions = []
    for row_number, row in enumerate(rows[1:], start=2):
        prompt = _text(row[0]) if row else ""
        if not prompt:
            continue

        answers = []
        answer_labels = []
        for column_index, value in enumerate(row[1:], start=1):
            if column_index - 1 < len(labels) and labels[column_index - 1] is None:
                continue
            answer = _text(value)
            if answer:
                answers.append(answer)
                answer_labels.append(
                    labels[column_index - 1]
                    if column_index - 1 < len(labels)
                    else f"Word {column_index}"
                )
        if not answers:
            continue

        questions.append({
            "id": f"{exercise_id}:row-{row_number}",
            "prompt": prompt,
            "answer": answers,
            "slot_labels": answer_labels,
            "group": re.sub(r"[.!?]+$", "", answers[-1]).casefold(),
        })

    if not questions:
        return None
    return {
        "id": exercise_id,
        "title": sheet.title,
        "description": f"Build questions from the {sheet.title} table.",
        "questions": questions,
    }


def _load_game_data(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        exercises = {}
        for sheet in workbook.worksheets:
            if not re.match(r"^level 1(?:\b|$)", sheet.title, re.I):
                continue
            exercise_id = f"level1-{_slug(sheet.title[7:])}"
            if exercise_id in exercises:
                exercise_id = f"{exercise_id}-{len(exercises) + 1}"
            exercise = _load_sheet_exercise(sheet, exercise_id)
            if exercise:
                exercises[exercise_id] = exercise

        if not exercises:
            raise ValueError(
                "Question game workbook has no usable sheets starting with 'level 1'."
            )
        return {"levels": exercises}
    finally:
        workbook.close()


def get_game_data():
    global _game_data_cache
    path = os.path.abspath(os.fspath(
        current_app.config.get("QUESTION_GAME_WORKBOOK", WORKBOOK_PATH)
    ))
    modified = os.path.getmtime(path)
    if (
        _game_data_cache is None
        or _game_data_cache["path"] != path
        or _game_data_cache["modified"] != modified
    ):
        _game_data_cache = {
            "path": path,
            "modified": modified,
            "data": _load_game_data(path),
        }
    return _game_data_cache["data"]


def _make_batches(questions):
    questions_by_group = {}
    for question in questions:
        questions_by_group.setdefault(question["group"], []).append(dict(question))

    batches = []
    for group_questions in questions_by_group.values():
        random.shuffle(group_questions)
        unique_words = {}
        for word in (
            word for question in group_questions for word in question["answer"]
        ):
            unique_words.setdefault(word.casefold(), word)
        tiles = [
            {"id": f"tile-{index}", "word": word}
            for index, word in enumerate(unique_words.values())
        ]
        random.shuffle(tiles)
        batches.append({
            "questions": [
                {
                    "id": question["id"],
                    "prompt": question["prompt"],
                    "slot_labels": question["slot_labels"],
                }
                for question in group_questions
            ],
            "word_bank": tiles,
        })
    random.shuffle(batches)
    return batches


@question_game_bp.route("/game")
def game_page():
    return render_template("games/question_game.html")


@question_game_bp.route("/api/exercises")
def get_exercises():
    return jsonify({
        "exercises": [
            {
                "id": exercise_id,
                "title": exercise["title"],
                "description": exercise["description"],
                "total": len(exercise["questions"]),
            }
            for exercise_id, exercise in get_game_data()["levels"].items()
        ]
    })


@question_game_bp.route("/api/round")
def get_round():
    level_id = request.args.get("level", "")
    level = get_game_data()["levels"].get(level_id)
    if level is None:
        return jsonify({"error": "Choose a valid question game exercise."}), 400
    return jsonify({
        "level": level_id,
        "title": level["title"],
        "total": len(level["questions"]),
        "batches": _make_batches(level["questions"]),
    })


@question_game_bp.route("/api/check", methods=["POST"])
def check_answers():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"error": "Submit answers as a JSON object."}), 400
    level_id = payload.get("level")
    level = get_game_data()["levels"].get(level_id)
    if level is None:
        return jsonify({"error": "Choose a valid question game exercise."}), 400
    answers = payload.get("answers")
    if not isinstance(answers, list) or not 1 <= len(answers) <= 3:
        return jsonify({"error": "Submit answers for one to three questions."}), 400

    questions_by_id = {
        question["id"]: question for question in level["questions"]
    }
    results = []
    seen_ids = set()
    for answer in answers:
        if not isinstance(answer, dict):
            return jsonify({"error": "Each answer must include a question and answer words."}), 400
        question_id = answer.get("id")
        words = answer.get("words")
        question = questions_by_id.get(question_id) if isinstance(question_id, str) else None
        if (
            question is None
            or question_id in seen_ids
            or not isinstance(words, list)
            or len(words) != len(question["answer"])
            or not all(isinstance(word, str) and word.strip() for word in words)
        ):
            return jsonify({"error": "Each answer must include a valid question and all answer words."}), 400
        seen_ids.add(question_id)
        correct = [word.strip().casefold() for word in words] == [
            word.casefold() for word in question["answer"]
        ]
        results.append({
            "id": question_id,
            "correct": correct,
            "correct_answer": question["answer"],
        })
    return jsonify({"results": results})
